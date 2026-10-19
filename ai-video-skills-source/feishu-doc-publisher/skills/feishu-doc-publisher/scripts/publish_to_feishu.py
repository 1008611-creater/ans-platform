#!/usr/bin/env python3
"""
One-click Markdown / DOCX to Feishu Online Document publisher.
Automatically converts Markdown (.md) + local images into a genuine Microsoft Word (.docx)
via python-docx, then imports it into Feishu Drive as a native Feishu Online Document
using the persisted QR-authenticated browser session (~/.gemini/feishu-browser-profile).
"""
import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
from docx import Document
from docx.shared import Inches

NODE_BIN = "/Users/eeo/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node"
WORKER_MJS = "/Users/eeo/.gemini/config/plugins/feishu-doc-publisher/skills/feishu-doc-publisher/scripts/feishu_playwright_worker.mjs"


def md_to_docx(md_path: str, out_docx_path: str) -> str:
    md_dir = os.path.dirname(os.path.abspath(md_path))
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    table_rows = []
    in_code = False
    code_lines = []
    tmp_img_dir = tempfile.mkdtemp(prefix="feishu_docx_imgs_")
    img_counter = 0

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        cols = max(len(r) for r in table_rows)
        t = doc.add_table(rows=len(table_rows), cols=cols)
        t.style = "Table Grid"
        for r_idx, r in enumerate(table_rows):
            for c_idx in range(len(r)):
                t.rows[r_idx].cells[c_idx].text = r[c_idx]
        table_rows = []

    for line in lines:
        s = line.strip()
        if s.startswith("```"):
            flush_table()
            if in_code:
                doc.add_paragraph("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue

        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip().replace("`", "").replace("**", "") for c in s.strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue
            table_rows.append(cells)
            continue
        else:
            flush_table()

        if not s or s == "---":
            continue

        m_img = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m_img:
            raw_path = m_img.group(2).replace("file://", "").strip()
            if not os.path.isabs(raw_path):
                raw_path = os.path.abspath(os.path.join(md_dir, raw_path))
            if os.path.exists(raw_path):
                img_counter += 1
                jpg_tmp = os.path.join(tmp_img_dir, f"img_{img_counter:03d}.jpg")
                subprocess.run(
                    ["sips", "-Z", "1280", "-s", "format", "jpeg", "-s", "formatOptions", "78", raw_path, "--out", jpg_tmp],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
                doc.add_picture(jpg_tmp if os.path.exists(jpg_tmp) else raw_path, width=Inches(6.0))
            continue

        clean = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1 (\2)", s).replace("`", "").replace("**", "")
        if s.startswith("# "):
            doc.add_heading(clean[2:].strip(), level=0)
        elif s.startswith("## "):
            doc.add_heading(clean[3:].strip(), level=1)
        elif s.startswith("### "):
            doc.add_heading(clean[4:].strip(), level=2)
        elif s.startswith("#### "):
            doc.add_heading(clean[5:].strip(), level=3)
        elif s.startswith("> "):
            doc.add_paragraph(clean[2:].strip())
        elif s.startswith("- "):
            doc.add_paragraph("• " + clean[2:].strip())
        else:
            doc.add_paragraph(clean)

    flush_table()
    doc.save(out_docx_path)
    return out_docx_path


def main():
    parser = argparse.ArgumentParser(description="Convert Markdown/DOCX and publish to Feishu Cloud Docs")
    parser.add_argument("input_file", help="Path to .md or .docx file")
    parser.add_argument("--drive-url", default="https://narcb9vct6.feishu.cn/drive/home/", help="Feishu Drive URL")
    args = parser.parse_args()

    input_path = os.path.abspath(args.input_file)
    if not os.path.exists(input_path):
        print(json.dumps({"ok": False, "error": f"File not found: {input_path}"}, ensure_ascii=False))
        sys.exit(1)

    if input_path.lower().endswith((".md", ".markdown")):
        docx_path = os.path.splitext(input_path)[0] + ".docx"
        md_to_docx(input_path, docx_path)
    else:
        docx_path = input_path

    payload = json.dumps({"file_path": docx_path, "drive_url": args.drive_url}, ensure_ascii=False)
    proc = subprocess.run([NODE_BIN, WORKER_MJS, "feishu_import_file_as_doc", payload], capture_output=True, text=True)
    if proc.stdout.strip():
        print(proc.stdout.strip())
    else:
        print(json.dumps({"ok": False, "stderr": proc.stderr}, ensure_ascii=False))


if __name__ == "__main__":
    main()
